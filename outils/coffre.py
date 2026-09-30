#!/usr/bin/env python3
"""Coffre des mots de passe du site Technologie (règle 8 de CLAUDE.md).

index.html est public : il ne contient aucun mot de passe ni code d'accès en
clair. Les mots de passe y figurent sous forme d'empreintes ; les valeurs en
clair utiles au professeur (impression des identifiants, codes à communiquer)
sont dans un coffre chiffré par le mot de passe professeur.

  python outils/coffre.py initialiser
      Convertit une version en clair (pwd, PMDP, code) : empreintes + coffre.
      Renouvelle les mots de passe des collègues, qui étaient publics.
  python outils/coffre.py ajouter --id lea.martin --prenom Léa --nom MARTIN --classe 4B
  python outils/coffre.py ajouter --id marc.dupont --prenom Marc --nom DUPONT --enseignant --classes 5B,5C
      Ajoute un compte ; le mot de passe généré s'affiche une seule fois.
  python outils/coffre.py renouveler --jusqu-au 2026-10-04
      Nouveaux mots de passe pour tous les élèves + planches Word à découper (dans
      Google Drive si présent). Les anciens restent acceptés jusqu'à la date donnée.
  python outils/coffre.py nettoyer
      Après cette date : retire les anciennes empreintes.

Le mot de passe professeur est demandé au clavier, jamais affiché ni stocké.
Le script s'arrête sans rien écrire si une ancre attendue manque.
"""
import argparse
import base64
import getpass
import hashlib
import json
import os
import re
import secrets
import sys

from cryptography.exceptions import InvalidTag
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

SEL = "stp-techno:"          # = EMPREINTE_SEL dans index.html
ITER_EMPREINTE = 100000      # = EMPREINTE_ITER dans index.html
ITER_COFFRE = 600000
RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INDEX = os.path.join(RACINE, "index.html")
OBJET = re.compile(r"\{[^{}]*\}")


class ErreurCoffre(Exception):
    pass


def _b64(octets):
    return base64.b64encode(octets).decode("ascii")


def empreinte(texte, sel):
    return _b64(hashlib.pbkdf2_hmac("sha256", texte.encode("utf-8"), sel.encode("utf-8"), ITER_EMPREINTE, 32))


def empreinte_mdp(ident, mdp):
    return empreinte(mdp, SEL + ident)


def empreinte_code(seq_id, code):
    return "#" + empreinte(code.upper(), SEL + "code:" + seq_id)


def _cle(mdp, sel, iterations):
    return hashlib.pbkdf2_hmac("sha256", mdp.encode("utf-8"), sel, iterations, 32)


def chiffrer(contenu, mdp):
    sel, iv = os.urandom(16), os.urandom(12)
    clair = json.dumps(contenu, ensure_ascii=False, sort_keys=True).encode("utf-8")
    data = AESGCM(_cle(mdp, sel, ITER_COFFRE)).encrypt(iv, clair, None)
    return {"v": 1, "it": ITER_COFFRE, "sel": _b64(sel), "iv": _b64(iv), "data": _b64(data)}


def dechiffrer(coffre, mdp):
    d = base64.b64decode
    try:
        clair = AESGCM(_cle(mdp, d(coffre["sel"]), coffre["it"])).decrypt(d(coffre["iv"]), d(coffre["data"]), None)
    except InvalidTag:
        raise ErreurCoffre("Mot de passe professeur incorrect : le coffre ne s'ouvre pas.")
    return json.loads(clair.decode("utf-8"))


def _syllabe():
    return secrets.choice("bdfgklmnprstvz") + secrets.choice("aeiou")


def nouveau_mdp():
    """Facile à dicter, sans lien avec le nom : « bako-rimu-47 »."""
    return _syllabe() + _syllabe() + "-" + _syllabe() + _syllabe() + "-" + "%02d" % secrets.randbelow(100)


def _bloc(texte, debut, fin="];"):
    """Positions du contenu compris entre la ligne `debut` et la première ligne `fin`."""
    m = re.search(re.escape(debut) + r"\r?\n(.*?)\r?\n" + re.escape(fin), texte, re.S)
    if not m:
        raise ErreurCoffre("Ancre introuvable dans index.html : " + debut)
    return m.start(1), m.end(1)


def _transformer(texte, debut, fonction, fin="];"):
    i, j = _bloc(texte, debut, fin)
    return texte[:i] + fonction(texte[i:j]) + texte[j:]


def _coffre_json(coffre):
    return json.dumps(coffre, separators=(",", ":"))


def initialiser(texte, mdp_prof):
    """Version en clair -> (texte converti, rapport). Ne touche qu'aux données."""
    if "COFFRE:" in texte:
        raise ErreurCoffre("Déjà fait : index.html contient déjà un coffre.")
    i, j = _bloc(texte, "var CFG = {", "};")
    ancien = re.search(r'PMDP:\s*"([^"]*)"', texte[i:j])
    if not ancien:
        raise ErreurCoffre("CFG.PMDP introuvable.")
    if len(mdp_prof) < 12:
        raise ErreurCoffre("Mot de passe professeur trop court : 12 caractères minimum.")
    if mdp_prof == ancien.group(1):
        raise ErreurCoffre("Le mot de passe professeur doit changer : l'ancien est public.")

    contenu = {"eleves": {}, "enseignants": {}, "codes": {}, "dcAnciens": []}
    renouveles = []

    def comptes(liste, renouveler):
        def convertir(m):
            obj = m.group(0)
            ident, pw = re.search(r'id:"([^"]+)"', obj), re.search(r'pwd:"([^"]*)"', obj)
            if not (ident and pw):
                raise ErreurCoffre("Compte sans identifiant ou sans mot de passe : " + obj[:60])
            mdp = pw.group(1)
            if renouveler:
                mdp = nouveau_mdp()
                renouveles.append(ident.group(1))
            contenu[liste][ident.group(1)] = mdp
            return obj[:pw.start()] + 'h:"' + empreinte_mdp(ident.group(1), mdp) + '"' + obj[pw.end():]
        return lambda bloc: OBJET.sub(convertir, bloc)

    def codes(bloc):
        def convertir(m):
            obj = m.group(0)
            code = re.search(r'code:"([^"]+)"', obj)
            if not code:
                return obj
            sid = re.search(r'id:"([^"]+)"', obj).group(1)
            contenu["codes"][sid] = code.group(1)
            return obj[:code.start()] + 'code:"' + empreinte_code(sid, code.group(1)) + '"' + obj[code.end():]
        return OBJET.sub(convertir, bloc)

    texte = _transformer(texte, "var ELEVES = [", comptes("eleves", False))
    texte = _transformer(texte, "var ENSEIGNANTS = [", comptes("enseignants", True))
    texte = _transformer(texte, "var SEQUENCES_DEF = [", codes)

    dc = re.search(r"var DC_CODES_ANCIENS = \[([^\]]*)\];", texte)
    if not dc:
        raise ErreurCoffre("Ancre introuvable dans index.html : var DC_CODES_ANCIENS")
    contenu["dcAnciens"] = re.findall(r'"([^"]*)"', dc.group(1))
    hachés = ", ".join('"' + empreinte_code("dc-corrige", c) + '"' for c in contenu["dcAnciens"])
    texte = texte[:dc.start(1)] + hachés + texte[dc.end(1):]

    coffre = chiffrer(contenu, mdp_prof)
    i, j = _bloc(texte, "var CFG = {", "};")
    cfg = re.sub(r'PMDP:\s*"[^"]*"', lambda m: "COFFRE: " + _coffre_json(coffre), texte[i:j])
    texte = texte[:i] + cfg + texte[j:]

    # Contrôles avant de rendre la main : rien en clair, coffre lisible
    for debut in ("var ELEVES = [", "var ENSEIGNANTS = ["):
        a, b = _bloc(texte, debut)
        if 'pwd:"' in texte[a:b]:
            raise ErreurCoffre("Un mot de passe en clair subsiste dans " + debut)
    i, j = _bloc(texte, "var CFG = {", "};")
    if "PMDP" in texte[i:j] or dechiffrer(coffre, mdp_prof) != contenu:
        raise ErreurCoffre("Contrôle final en échec : rien n'a été écrit.")
    rapport = {"eleves": len(contenu["eleves"]), "enseignants": len(contenu["enseignants"]),
               "codes": len(contenu["codes"]), "anciens": len(contenu["dcAnciens"]), "renouveles": renouveles}
    return texte, rapport


def ajouter(texte, mdp_prof, compte):
    """Ajoute un élève (ou un collègue si compte["enseignant"]). -> (texte, mot de passe)."""
    m = re.search(r"COFFRE: (\{[^\r\n]*\})", texte)
    if not m:
        raise ErreurCoffre("Pas de coffre : lancer d'abord « initialiser ».")
    contenu = dechiffrer(json.loads(m.group(1)), mdp_prof)
    ident = compte["id"].strip().lower()
    if not re.fullmatch(r"[a-z0-9.-]+", ident):
        raise ErreurCoffre("Identifiant invalide (lettres sans accent, chiffres, point, tiret) : " + ident)
    if re.search(r'id:"' + re.escape(ident) + '"', texte):
        raise ErreurCoffre("Identifiant déjà utilisé : " + ident)
    for cle in ("p", "n", "c"):
        if '"' in compte.get(cle, ""):
            raise ErreurCoffre("Guillemet interdit dans les champs.")
    mdp = nouveau_mdp()
    h = empreinte_mdp(ident, mdp)
    if compte.get("enseignant"):
        classes = ",".join('"' + c + '"' for c in compte.get("classes", []))
        ligne, debut, liste = '  { id:"%s", p:"%s", n:"%s", h:"%s", classes:[%s] }' % (ident, compte["p"], compte["n"], h, classes), "var ENSEIGNANTS = [", "enseignants"
    else:
        c = compte["c"]
        ligne, debut, liste = '  {id:"%s", p:"%s", n:"%s", c:"%s", niveau:"%sème", h:"%s"},' % (ident, compte["p"], compte["n"], c, c[:1], h), "var ELEVES = [", "eleves"
    nl = "\r\n" if "\r\n" in texte else "\n"

    def inserer(bloc):
        k = bloc.rfind("}")
        if k >= 0 and not bloc[k + 1:].lstrip().startswith(","):
            bloc = bloc[:k + 1] + "," + bloc[k + 1:]
        return bloc + nl + ligne

    texte = _transformer(texte, debut, inserer)
    contenu[liste][ident] = mdp
    m = re.search(r"COFFRE: (\{[^\r\n]*\})", texte)
    texte = texte[:m.start(1)] + _coffre_json(chiffrer(contenu, mdp_prof)) + texte[m.end(1):]
    return texte, mdp


def _champ(obj, nom):
    m = re.search(nom + r':"([^"]*)"', obj)
    return m.group(1) if m else ""


def renouveler(texte, mdp_prof, jusqu_au):
    """Nouveaux mots de passe pour tous les comptes élèves. L'ancienne empreinte passe
    en h0 : index.html l'accepte encore jusqu'à la date `jusqu_au` (AAAA-MM-JJ), le
    temps de distribuer les planches. -> (texte, comptes pour les planches)."""
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", jusqu_au or ""):
        raise ErreurCoffre("Date de fin des anciens mots de passe attendue au format AAAA-MM-JJ.")
    m = re.search(r"COFFRE: (\{[^\r\n]*\})", texte)
    if not m:
        raise ErreurCoffre("Pas de coffre : lancer d'abord « initialiser ».")
    contenu = dechiffrer(json.loads(m.group(1)), mdp_prof)
    a, b = _bloc(texte, "var ELEVES = [")
    if 'h0:"' in texte[a:b]:
        raise ErreurCoffre("Un renouvellement est déjà en cours : lancer « nettoyer » après sa date, puis recommencer.")
    comptes = []

    def convertir(mo):
        obj = mo.group(0)
        ident, h = re.search(r'id:"([^"]+)"', obj), re.search(r'h:"([^"]*)"', obj)
        if not (ident and h):
            raise ErreurCoffre("Compte sans empreinte : " + obj[:60])
        mdp = nouveau_mdp()
        contenu["eleves"][ident.group(1)] = mdp
        comptes.append({"id": ident.group(1), "p": _champ(obj, "p"), "n": _champ(obj, "n"),
                        "c": _champ(obj, "c"), "role": _champ(obj, "role"), "mdp": mdp})
        return obj[:h.start()] + 'h:"' + empreinte_mdp(ident.group(1), mdp) + '", h0:"' + h.group(1) + '"' + obj[h.end():]

    texte = _transformer(texte, "var ELEVES = [", lambda bloc: OBJET.sub(convertir, bloc))
    nl = "\r\n" if "\r\n" in texte else "\n"
    m = re.search(r"(\r?\n)([ \t]*)COFFRE: \{[^\r\n]*\}", texte)
    nouveau = (m.group(1) + m.group(2) + 'ANCIENS_JUSQUAU: "' + jusqu_au + '",' + nl + m.group(2)
               + "COFFRE: " + _coffre_json(chiffrer(contenu, mdp_prof)))
    return texte[:m.start()] + nouveau + texte[m.end():], comptes


def nettoyer(texte):
    """Après la date de transition : retire les anciennes empreintes (h0) et la date."""
    texte = _transformer(texte, "var ELEVES = [", lambda bloc: re.sub(r', h0:"[^"]*"', "", bloc))
    return re.sub(r'\r?\n[ \t]*ANCIENS_JUSQUAU: "[^"]*",', "", texte)


ADRESSE_SITE = "blanchardromain-cyber.github.io/technologie-stpierre"


def planches(comptes, chemin):
    """Planches A4 à découper, une série par classe : 2 étiquettes par ligne, traits
    pointillés. Écrit le fichier Word sans rien afficher."""
    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        s.top_margin = s.bottom_margin = Cm(1.2)
        s.left_margin = s.right_margin = Cm(1.2)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(11)

    def pointilles(cellule):
        bords = OxmlElement("w:tcBorders")
        for cote in ("top", "left", "bottom", "right"):
            e = OxmlElement("w:" + cote)
            e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "8A94A6")
            bords.append(e)
        cellule._tc.get_or_add_tcPr().append(bords)

    def ligne(cellule, texte, taille, gras=False, couleur=None, premiere=False):
        par = cellule.paragraphs[0] if premiere else cellule.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        if premiere:
            par.paragraph_format.space_before = Pt(5)
        r = par.add_run(texte)
        r.font.size, r.bold = Pt(taille), gras
        if couleur:
            r.font.color.rgb = RGBColor.from_string(couleur)

    groupes = {}
    for c in comptes:
        groupes.setdefault("Autres comptes" if c.get("role") else "Classe " + (c["c"] or "?"), []).append(c)
    ordre = sorted(k for k in groupes if k != "Autres comptes") + (["Autres comptes"] if "Autres comptes" in groupes else [])
    for g, nom in enumerate(ordre):
        liste = sorted(groupes[nom], key=lambda c: (c["n"].upper(), c["p"]))
        titre = doc.add_paragraph()
        r = titre.add_run(nom + " — identifiants du site de Technologie (à découper)")
        r.bold, r.font.size = True, Pt(13)
        table = doc.add_table(rows=(len(liste) + 1) // 2, cols=2)
        table.autofit = False
        for i, c in enumerate(liste):
            rangee = table.rows[i // 2]
            rangee.height, rangee.height_rule = Cm(3.6), WD_ROW_HEIGHT_RULE.AT_LEAST   # 7 lignes, 14 élèves par feuille
            cellule = rangee.cells[i % 2]
            cellule.width = Cm(9.3)
            pointilles(cellule)
            ligne(cellule, c["p"] + " " + c["n"] + ("   —   " + c["c"] if c["c"] else ""), 13, gras=True, premiere=True)
            ligne(cellule, "Identifiant :  " + c["id"], 12)
            ligne(cellule, "Mot de passe :  " + c["mdp"], 15, gras=True)
            ligne(cellule, "Site : " + ADRESSE_SITE, 9, couleur="5A6475")
        for cellule in table.rows[-1].cells:
            pointilles(cellule)
        if g < len(ordre) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(chemin)


def _saisie_masquee(question):
    """Petite fenêtre de saisie masquée. Elle fonctionne même quand le terminal ne
    transmet pas la frappe à getpass (terminal intégré, Python embarqué de WAPT).
    Repli sur getpass si tkinter manque."""
    try:
        import tkinter as tk
    except ImportError:
        return getpass.getpass(question + " ")
    racine = tk.Tk()
    racine.title("Coffre du site Technologie")
    racine.attributes("-topmost", True)
    reponse = {"texte": None}
    tk.Label(racine, text=question, padx=18, pady=12).pack()
    champ = tk.Entry(racine, show="•", width=38)
    champ.pack(padx=18)
    champ.focus_force()

    def valider(_evenement=None):
        reponse["texte"] = champ.get()
        racine.destroy()

    champ.bind("<Return>", valider)
    tk.Button(racine, text="Valider", command=valider).pack(pady=12)
    racine.mainloop()
    if reponse["texte"] is None:
        raise ErreurCoffre("Saisie annulée.")
    return reponse["texte"]


def _mot_de_passe_prof(args, nouveau):
    if args.mdp_env:
        return os.environ[args.mdp_env]
    mdp = _saisie_masquee("Nouveau mot de passe professeur (12 caractères minimum) :" if nouveau else "Mot de passe professeur :")
    if nouveau and _saisie_masquee("Confirmez le nouveau mot de passe :") != mdp:
        raise ErreurCoffre("Les deux saisies diffèrent.")
    return mdp


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    p = argparse.ArgumentParser(description="Coffre des mots de passe du site Technologie.")
    sous = p.add_subparsers(dest="commande", required=True)
    for nom in ("initialiser", "ajouter", "renouveler", "nettoyer"):
        s = sous.add_parser(nom)
        s.add_argument("--fichier", default=INDEX)
        s.add_argument("--sortie", help="par défaut, réécrit --fichier")
        s.add_argument("--mdp-env", help="TESTS SEULEMENT : variable d'environnement contenant le mot de passe professeur")
        if nom == "renouveler":
            s.add_argument("--jusqu-au", required=True, help="dernier jour où l'ancien mot de passe est accepté (AAAA-MM-JJ)")
            s.add_argument("--planches", default=os.path.join("G:\\Mon Drive" if os.path.isdir("G:\\Mon Drive") else RACINE,
                                                              "Identifiants techno — planches.docx"),
                           help="fichier Word des planches à découper")
        if nom == "ajouter":
            s.add_argument("--id", required=True)
            s.add_argument("--prenom", required=True)
            s.add_argument("--nom", required=True)
            s.add_argument("--classe", default="")
            s.add_argument("--enseignant", action="store_true")
            s.add_argument("--classes", default="", help="collègue : classes suivies, séparées par des virgules")
    a = p.parse_args(argv)
    sortie = a.sortie or a.fichier
    if a.mdp_env and os.path.abspath(sortie) == os.path.abspath(INDEX):
        p.error("--mdp-env sert aux tests : il n'écrit jamais sur le index.html du site.")
    if a.commande == "ajouter" and not a.enseignant and not a.classe:
        p.error("--classe est obligatoire pour un élève.")
    try:
        with open(a.fichier, encoding="utf-8", newline="") as f:
            texte = f.read()
        if a.commande == "initialiser":
            texte, r = initialiser(texte, _mot_de_passe_prof(a, True))
            message = ("Conversion faite : %d élèves, %d collègue(s), %d codes, %d anciens codes du corrigé DC.\n"
                       "Mots de passe renouvelés (visibles dans le site, onglet Config) : %s"
                       % (r["eleves"], r["enseignants"], r["codes"], r["anciens"], ", ".join(r["renouveles"]) or "aucun"))
        elif a.commande == "renouveler":
            texte, comptes = renouveler(texte, _mot_de_passe_prof(a, False), a.jusqu_au)
            planches(comptes, a.planches)
            message = ("Nouveaux mots de passe : %d comptes. Les anciens restent acceptés jusqu'au %s inclus.\n"
                       "Planches à découper : %s" % (len(comptes), a.jusqu_au, a.planches))
        elif a.commande == "nettoyer":
            texte = nettoyer(texte)
            message = "Anciennes empreintes retirées."
        else:
            compte = {"id": a.id, "p": a.prenom, "n": a.nom, "c": a.classe, "enseignant": a.enseignant,
                      "classes": [c.strip() for c in a.classes.split(",") if c.strip()]}
            texte, mdp = ajouter(texte, _mot_de_passe_prof(a, False), compte)
            message = "Compte ajouté : %s — mot de passe à lui transmettre : %s (aussi visible dans Config)" % (a.id.lower(), mdp)
    except ErreurCoffre as e:
        print("Erreur : " + str(e) + " Rien n'a été écrit.", file=sys.stderr)
        return 1
    with open(sortie, "w", encoding="utf-8", newline="") as f:
        f.write(texte)
    print(message)
    print("Fichier écrit : " + sortie)
    return 0


if __name__ == "__main__":
    sys.exit(main())
