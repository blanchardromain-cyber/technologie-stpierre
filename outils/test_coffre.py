"""Tests de outils/coffre.py — données entièrement fictives.

Lancer : python outils/test_coffre.py
"""
import base64
import hashlib
import json
import os
import re
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import coffre  # noqa: E402

MDP_PROF = "phrase-de-test-assez-longue"

FIXTURE = (
    'var CFG = {\r\n'
    '  PLOGIN: "prof.test",\r\n'
    '  PMDP: "ancien-mdp-prof"\r\n'
    '};\r\n'
    'var ELEVES = [\r\n'
    '  /* commentaire */\r\n'
    '  {id:"lea.test", p:"Léa", n:"TEST", c:"4B", niveau:"4ème", pwd:"lea1234"},\r\n'
    '  {id:"tom.test", p:"Tom", n:"TEST", c:"5A", niveau:"5ème", pwd:"tom5678", role:"test"}\r\n'
    '];\r\n'
    'var ENSEIGNANTS = [\r\n'
    '  { id:"ana.test", p:"Ana", n:"TEST", pwd:"ana0000", classes:["5A","5B"] }\r\n'
    '];\r\n'
    'var SEQUENCES_DEF = [\r\n'
    '  {id:"s1", ordre:1, titre:"Un", code:"abc26", open:true},\r\n'
    '  {id:"s2", ordre:2, titre:"Deux", code:"", open:true},\r\n'
    '  {id:"dc-corrige", ordre:3, titre:"DC", code:"CODX00", open:true}\r\n'
    '];\r\n'
    'var DC_CODES_ANCIENS = ["VIEUX2026", "CODX00"];\r\n'
)


def lire_coffre(texte):
    m = re.search(r"COFFRE: (\{[^\r\n]*\})", texte)
    return json.loads(m.group(1))


def champ(texte, ident, nom):
    m = re.search(r'\{[^{}]*id:"' + re.escape(ident) + r'"[^{}]*\}', texte)
    return re.search(nom + r':"([^"]*)"', m.group(0)).group(1)


class TestEmpreintes(unittest.TestCase):
    def test_empreinte_conforme_a_pbkdf2(self):
        attendu = base64.b64encode(hashlib.pbkdf2_hmac("sha256", b"abc", b"sel", 100000, 32)).decode()
        self.assertEqual(coffre.empreinte("abc", "sel"), attendu)

    def test_empreinte_code_en_majuscules(self):
        self.assertEqual(coffre.empreinte_code("s1", "abc26"), coffre.empreinte_code("s1", "ABC26"))
        self.assertTrue(coffre.empreinte_code("s1", "abc26").startswith("#"))

    def test_nouveau_mdp(self):
        m = coffre.nouveau_mdp()
        self.assertRegex(m, r"^[a-z]{4}-[a-z]{4}-\d{2}$")
        self.assertNotEqual(m, coffre.nouveau_mdp())


class TestInitialiser(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.texte, cls.rapport = coffre.initialiser(FIXTURE, MDP_PROF)

    def test_plus_aucun_secret_en_clair(self):
        for secret in ("pwd:", "PMDP", "lea1234", "tom5678", "ana0000", "ancien-mdp-prof", "abc26", "VIEUX2026", "CODX00"):
            self.assertNotIn(secret, self.texte, secret)

    def test_empreintes_des_eleves(self):
        self.assertEqual(champ(self.texte, "lea.test", "h"), coffre.empreinte_mdp("lea.test", "lea1234"))
        self.assertEqual(champ(self.texte, "tom.test", "h"), coffre.empreinte_mdp("tom.test", "tom5678"))
        self.assertEqual(champ(self.texte, "tom.test", "role"), "test")

    def test_coffre_complet_et_collegue_renouvele(self):
        contenu = coffre.dechiffrer(lire_coffre(self.texte), MDP_PROF)
        self.assertEqual(contenu["eleves"], {"lea.test": "lea1234", "tom.test": "tom5678"})
        self.assertEqual(contenu["codes"], {"s1": "abc26", "dc-corrige": "CODX00"})
        self.assertEqual(contenu["dcAnciens"], ["VIEUX2026", "CODX00"])
        nouveau = contenu["enseignants"]["ana.test"]
        self.assertNotEqual(nouveau, "ana0000")
        self.assertEqual(champ(self.texte, "ana.test", "h"), coffre.empreinte_mdp("ana.test", nouveau))
        self.assertEqual(self.rapport["renouveles"], ["ana.test"])

    def test_codes(self):
        self.assertEqual(champ(self.texte, "s1", "code"), coffre.empreinte_code("s1", "abc26"))
        self.assertEqual(champ(self.texte, "s2", "code"), "")
        anciens = re.search(r"var DC_CODES_ANCIENS = \[([^\]]*)\];", self.texte).group(1)
        self.assertIn(coffre.empreinte_code("dc-corrige", "VIEUX2026"), anciens)

    def test_mauvais_mot_de_passe_refuse(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.dechiffrer(lire_coffre(self.texte), "autre-phrase-de-test")

    def test_fins_de_ligne_conservees(self):
        self.assertEqual(self.texte.count("\r\n"), FIXTURE.count("\r\n"))
        self.assertNotIn("\n", self.texte.replace("\r\n", ""))

    def test_rapport(self):
        self.assertEqual((self.rapport["eleves"], self.rapport["enseignants"], self.rapport["codes"], self.rapport["anciens"]), (2, 1, 2, 2))

    def test_deja_initialise(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.initialiser(self.texte, MDP_PROF)


class TestRefus(unittest.TestCase):
    def test_ancre_manquante(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.initialiser(FIXTURE.replace("var DC_CODES_ANCIENS", "var AUTRE"), MDP_PROF)

    def test_mot_de_passe_prof_trop_court(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.initialiser(FIXTURE, "court")

    def test_mot_de_passe_prof_inchange(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.initialiser(FIXTURE.replace("ancien-mdp-prof", MDP_PROF), MDP_PROF)

    def test_mode_test_n_ecrit_pas_sur_le_site(self):
        os.environ["COFFRE_MDP_TEST_UNITAIRE"] = MDP_PROF
        with self.assertRaises(SystemExit):
            coffre.main(["initialiser", "--mdp-env", "COFFRE_MDP_TEST_UNITAIRE"])


class TestAjouter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base, _ = coffre.initialiser(FIXTURE, MDP_PROF)

    def test_ajouter_eleve(self):
        texte, mdp = coffre.ajouter(self.base, MDP_PROF, {"id": "zoe.test", "p": "Zoé", "n": "TEST", "c": "4C"})
        self.assertNotIn(mdp, texte)
        self.assertEqual(champ(texte, "zoe.test", "h"), coffre.empreinte_mdp("zoe.test", mdp))
        self.assertEqual(champ(texte, "zoe.test", "niveau"), "4ème")
        self.assertEqual(coffre.dechiffrer(lire_coffre(texte), MDP_PROF)["eleves"]["zoe.test"], mdp)
        self.assertIn('role:"test"},\r\n  {id:"zoe.test"', texte)
        self.assertNotIn("\n", texte.replace("\r\n", ""))

    def test_ajouter_collegue(self):
        texte, mdp = coffre.ajouter(self.base, MDP_PROF, {"id": "bob.test", "p": "Bob", "n": "TEST", "enseignant": True, "classes": ["5A"]})
        self.assertIn('] },\r\n  { id:"bob.test"', texte)
        self.assertEqual(coffre.dechiffrer(lire_coffre(texte), MDP_PROF)["enseignants"]["bob.test"], mdp)

    def test_identifiant_existant_refuse(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.ajouter(self.base, MDP_PROF, {"id": "lea.test", "p": "Léa", "n": "TEST", "c": "4B"})

    def test_mauvais_mot_de_passe_prof(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.ajouter(self.base, "autre-phrase-de-test", {"id": "zoe.test", "p": "Zoé", "n": "TEST", "c": "4C"})


class TestCode(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base, _ = coffre.initialiser(FIXTURE, MDP_PROF)
        cls.texte, cls.code = coffre.poser_code(cls.base, MDP_PROF, "s2")

    def test_code_genere(self):
        self.assertRegex(self.code, r"^[A-Z]{6}\d{2}$")
        self.assertNotEqual(coffre.nouveau_code(), coffre.nouveau_code())

    def test_empreinte_dans_sequences_def(self):
        self.assertEqual(champ(self.texte, "s2", "code"), coffre.empreinte_code("s2", self.code))

    def test_code_dans_le_coffre_et_jamais_en_clair(self):
        contenu = coffre.dechiffrer(lire_coffre(self.texte), MDP_PROF)
        self.assertEqual(contenu["codes"]["s2"], self.code)
        self.assertNotIn(self.code, self.texte.upper())

    def test_autres_sequences_intactes(self):
        self.assertEqual(champ(self.texte, "s1", "code"), champ(self.base, "s1", "code"))
        contenu = coffre.dechiffrer(lire_coffre(self.texte), MDP_PROF)
        self.assertEqual(contenu["codes"]["s1"], "abc26")

    def test_code_choisi_mis_en_majuscules(self):
        texte, code = coffre.poser_code(self.base, MDP_PROF, "s2", "zorba77")
        self.assertEqual(code, "ZORBA77")
        self.assertEqual(champ(texte, "s2", "code"), coffre.empreinte_code("s2", "ZORBA77"))

    def test_code_devinable_refuse(self):
        # « SEQUENCES » est un mot du site ; « DCCORRIGE » reprend le nom de la séquence.
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.poser_code(self.base, MDP_PROF, "s2", "sequences")
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.poser_code(self.base, MDP_PROF, "dc-corrige", "dccorrige")

    def test_code_invalide_refuse(self):
        for mauvais in ("AB12", "ÉCOLE12", "AB-1234"):
            with self.assertRaises(coffre.ErreurCoffre):
                coffre.poser_code(self.base, MDP_PROF, "s2", mauvais)

    def test_sequence_inconnue_refusee(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.poser_code(self.base, MDP_PROF, "s9")

    def test_mauvais_mot_de_passe_prof(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.poser_code(self.base, "pas-le-bon-mot-de-passe", "s2")

    def test_crlf(self):
        self.assertEqual(self.texte.count("\r\n"), self.base.count("\r\n"))


class TestRenouveler(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base, _ = coffre.initialiser(FIXTURE, MDP_PROF)
        cls.texte, cls.comptes = coffre.renouveler(cls.base, MDP_PROF, "2026-10-04")

    def test_nouveaux_mots_de_passe_et_ancienne_empreinte_gardee(self):
        contenu = coffre.dechiffrer(lire_coffre(self.texte), MDP_PROF)
        for ident, ancien in (("lea.test", "lea1234"), ("tom.test", "tom5678")):
            nouveau = contenu["eleves"][ident]
            self.assertNotEqual(nouveau, ancien)
            self.assertRegex(nouveau, r"^[a-z]{4}-[a-z]{4}-\d{2}$")
            self.assertEqual(champ(self.texte, ident, "h"), coffre.empreinte_mdp(ident, nouveau))
            self.assertEqual(champ(self.texte, ident, "h0"), coffre.empreinte_mdp(ident, ancien))

    def test_date_de_transition_dans_cfg(self):
        self.assertIn('ANCIENS_JUSQUAU: "2026-10-04",', self.texte)
        self.assertIsNotNone(lire_coffre(self.texte))

    def test_comptes_pour_les_planches(self):
        par_id = {c["id"]: c for c in self.comptes}
        self.assertEqual(set(par_id), {"lea.test", "tom.test"})
        self.assertEqual((par_id["lea.test"]["p"], par_id["lea.test"]["c"], par_id["tom.test"]["role"]), ("Léa", "4B", "test"))
        contenu = coffre.dechiffrer(lire_coffre(self.texte), MDP_PROF)
        self.assertEqual(par_id["lea.test"]["mdp"], contenu["eleves"]["lea.test"])

    def test_collegues_et_codes_intacts(self):
        self.assertEqual(champ(self.texte, "ana.test", "h"), champ(self.base, "ana.test", "h"))
        self.assertEqual(champ(self.texte, "s1", "code"), champ(self.base, "s1", "code"))

    def test_crlf(self):
        self.assertNotIn("\n", self.texte.replace("\r\n", ""))

    def test_deux_transitions_refusees(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.renouveler(self.texte, MDP_PROF, "2026-10-04")

    def test_date_invalide(self):
        with self.assertRaises(coffre.ErreurCoffre):
            coffre.renouveler(self.base, MDP_PROF, "4 octobre")

    def test_nettoyer(self):
        propre = coffre.nettoyer(self.texte)
        self.assertNotIn('h0:"', propre)
        self.assertNotIn("ANCIENS_JUSQUAU", propre)
        self.assertEqual(champ(propre, "lea.test", "h"), champ(self.texte, "lea.test", "h"))
        self.assertNotIn("\n", propre.replace("\r\n", ""))

    def test_planches_docx(self):
        import tempfile
        from docx import Document
        chemin = os.path.join(tempfile.mkdtemp(), "planches.docx")
        coffre.planches(self.comptes, chemin)
        doc = Document(chemin)
        texte = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for c in self.comptes:
            self.assertIn(c["id"], texte)
            self.assertIn(c["mdp"], texte)
        self.assertIn("Classe 4B", "\n".join(p.text for p in doc.paragraphs))
        self.assertIn("Autres comptes", "\n".join(p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main(verbosity=1)
